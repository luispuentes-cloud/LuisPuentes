from __future__ import annotations

import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from xllib.inspect import load_workbook
from xllib.lint.api import lint
from xllib.lint.config import Config
from xllib.lint.registry import registered_rules

_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _rule(rule_id: str):
    return (next(rule for rule in registered_rules() if rule.id == rule_id),)


def _report(path: Path, rule_id: str):
    return lint(load_workbook(path), Config(), _rule(rule_id))


def _save(book: Workbook, path: Path) -> Path:
    book.save(path)
    return path


def _set_cache(
    path: Path, sheet_number: int, coordinate: str, value: str, cell_type: str | None = None
) -> None:
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    member = f"xl/worksheets/sheet{sheet_number}.xml"
    root = ET.fromstring(members[member])
    cell = root.find(f".//{{{_NS}}}c[@r='{coordinate}']")
    assert cell is not None
    if cell_type is not None:
        cell.set("t", cell_type)
    cached = cell.find(f"{{{_NS}}}v")
    if cached is None:
        cached = ET.SubElement(cell, f"{{{_NS}}}v")
    cached.text = value
    members[member] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)


def test_xl001_unlabelled_named_constant(tmp_path: Path) -> None:
    book = Workbook()
    inputs = book.active
    inputs.title = "Inputs"
    inputs["C4"] = 0.85
    calc = book.create_sheet("Calc")
    calc["D7"] = "=FTE_TARGET"
    book.defined_names.add(DefinedName("FTE_TARGET", attr_text="'Inputs'!$C$4"))
    report = _report(_save(book, tmp_path / "xl001.xlsx"), "XL001")
    assert [item.rule_id for item in report.findings] == ["XL001"]
    assert "FTE_TARGET" in report.findings[0].message


def test_xl002_numeric_literal_with_positional_exemption(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet["A1"] = 10
    sheet["B1"] = "=A1*0.85"
    sheet["B2"] = "=ROUND(A1,2)"
    config = Config(
        rules={"XL002": {"positional_exemptions": {"ROUND": [2]}}},
    )
    path = _save(book, tmp_path / "xl002.xlsx")
    report = lint(load_workbook(path), config, _rule("XL002"))
    assert [item.evidence["literal"] for item in report.findings] == ["0.85"]


def test_xl003_mixed_formula_shape(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet["D9"] = "=D8*$C$4"
    sheet["E9"] = "=E8*$C$4"
    sheet["F9"] = "=SUM(D9:E9)"
    report = _report(_save(book, tmp_path / "xl003.xlsx"), "XL003")
    assert len(report.findings) == 1
    assert report.findings[0].evidence["inferred_run"] == "D9:F9"


def test_xl004_text_in_numeric_cell(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet["E12"] = "N/A - no FTE load"
    sheet["E12"].number_format = "$#,##0"
    report = _report(_save(book, tmp_path / "xl004.xlsx"), "XL004")
    assert [item.locus.ref for item in report.findings] == ["E12"]


def test_xl005_text_in_aggregate_range(tmp_path: Path) -> None:
    book = Workbook()
    calc = book.active
    calc.title = "Calc"
    calc["E10"] = 1
    calc["E11"] = "not applicable"
    summary = book.create_sheet("Summary")
    summary["D8"] = "=SUM(Calc!E10:E11)"
    path = _save(book, tmp_path / "xl005.xlsx")
    _set_cache(path, 2, "D8", "1")
    report = _report(path, "XL005")
    assert [item.locus.ref for item in report.findings] == ["E11"]
    assert report.findings[0].evidence["aggregate"] == "Summary!D8"


def test_xl004_names_a_cached_error_rather_than_text(tmp_path: Path) -> None:
    """An error is a broken formula, not a formatting slip, and must say so."""
    book = Workbook()
    sheet = book.active
    sheet["E12"] = "=1/0"
    sheet["E12"].number_format = "$#,##0"
    path = _save(book, tmp_path / "xl004-error.xlsx")
    _set_cache(path, 1, "E12", "#DIV/0!", cell_type="e")
    report = _report(path, "XL004")
    assert [item.locus.ref for item in report.findings] == ["E12"]
    finding = report.findings[0]
    assert finding.evidence["kind"] == "error"
    assert "#DIV/0!" in finding.message
    assert "text" not in finding.message
    assert "status column" not in finding.remediation


def test_xl005_names_a_cached_error_rather_than_text(tmp_path: Path) -> None:
    book = Workbook()
    calc = book.active
    calc.title = "Calc"
    calc["E10"] = 1
    calc["E11"] = "=1/0"
    summary = book.create_sheet("Summary")
    summary["D8"] = "=SUM(Calc!E10:E11)"
    path = _save(book, tmp_path / "xl005-error.xlsx")
    _set_cache(path, 1, "E11", "#DIV/0!", cell_type="e")
    _set_cache(path, 2, "D8", "1")
    report = _report(path, "XL005")
    assert [item.locus.ref for item in report.findings] == ["E11"]
    finding = report.findings[0]
    assert finding.evidence["kind"] == "error"
    assert "#DIV/0!" in finding.message
    assert "text" not in finding.message


def test_xl004_still_calls_a_plain_string_text(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet["E12"] = "N/A - no FTE load"
    sheet["E12"].number_format = "$#,##0"
    report = _report(_save(book, tmp_path / "xl004-text.xlsx"), "XL004")
    # Behaviour before shape: asserting the evidence key first made this fail
    # on a KeyError and never reach the message it exists to protect.
    assert report.findings[0].message == "text value in a numeric-formatted cell"
    assert report.findings[0].evidence["kind"] == "text"


def test_xl004_does_not_call_error_shaped_text_an_error(tmp_path: Path) -> None:
    """Only the file's own type system may declare a cell an error."""
    book = Workbook()
    sheet = book.active
    sheet["E12"] = " #n/a "
    sheet["E12"].number_format = "$#,##0"
    report = _report(_save(book, tmp_path / "xl004-lookalike.xlsx"), "XL004")
    finding = report.findings[0]
    assert finding.evidence["kind"] == "text"
    assert "repair the formula" not in finding.remediation


def test_xl103_fails_a_check_evaluating_to_false(tmp_path: Path) -> None:
    """`False == 0` is true in Python; a failing boolean check must not pass."""
    book = Workbook()
    sheet = book.active
    sheet["A1"] = False
    book.defined_names.add(DefinedName("chk_bool_tie", attr_text="'Sheet'!$A$1"))
    report = _report(_save(book, tmp_path / "xl103-false.xlsx"), "XL103")
    assert len(report.findings) == 1
    assert "False" in report.findings[0].message


def test_xl103_passes_a_check_evaluating_to_true(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet["A1"] = True
    book.defined_names.add(DefinedName("chk_bool_tie", attr_text="'Sheet'!$A$1"))
    report = _report(_save(book, tmp_path / "xl103-true.xlsx"), "XL103")
    assert report.findings == ()


def test_xl006_inconsistent_period_axis(tmp_path: Path) -> None:
    book = Workbook()
    calc = book.active
    calc.title = "Calc"
    summary = book.create_sheet("Summary")
    for column, label in enumerate(("FY26", "FY27", "FY28"), start=4):
        calc.cell(3, column, label)
        summary.cell(3, column + 1, label)
    report = _report(_save(book, tmp_path / "xl006.xlsx"), "XL006")
    assert len(report.findings) == 1


def test_xl101_estimated_section_count(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    for row in (1, 3, 5, 7, 9, 11):
        sheet.cell(row, 1, f"Section {row}")
    report = _report(_save(book, tmp_path / "xl101.xlsx"), "XL101")
    assert len(report.findings) == 1
    assert report.findings[0].severity.value == "WARN"


def test_xl102_no_check_convention(tmp_path: Path) -> None:
    book = Workbook()
    report = _report(_save(book, tmp_path / "xl102.xlsx"), "XL102")
    assert len(report.findings) == 1


def test_xl103_check_not_passing(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet["A1"] = 1450
    book.defined_names.add(DefinedName("chk_total_tie", attr_text="'Sheet'!$A$1"))
    report = _report(_save(book, tmp_path / "xl103.xlsx"), "XL103")
    assert len(report.findings) == 1
    assert "1450" in report.findings[0].message


@pytest.mark.parametrize(
    ("style", "expected"),
    [
        ("none", 1),
        ("number", 1),
        ("font", 1),
        ("fill", 1),
        ("both", 0),
    ],
)
def test_xl104_requires_font_and_fill(tmp_path: Path, style: str, expected: int) -> None:
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet["A1"] = 5
    sheet["B1"] = "=A1"
    if style in {"number", "both"}:
        sheet["A1"].number_format = "0.00"
    if style in {"font", "both"}:
        sheet["A1"].font = Font(color="0000FF")
    if style in {"fill", "both"}:
        sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="DCE6F1")
    report = _report(_save(book, tmp_path / f"xl104-{style}.xlsx"), "XL104")
    assert len(report.findings) == expected


def test_xl006_info_when_no_period_axis_detected(tmp_path: Path) -> None:
    book = Workbook()
    book.active["A1"] = "not a period"
    report = _report(_save(book, tmp_path / "xl006-info.xlsx"), "XL006")
    assert len(report.findings) == 1
    assert report.findings[0].status.value == "INFO"
