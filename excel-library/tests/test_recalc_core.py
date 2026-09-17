from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from xllib.recalc import Availability, RecalcFailed, recalculate
from xllib.recalc.backends import FormulasBackend
from xllib.recalc.backends import formulas as formulas_backend
from xllib.recalc.cache import CacheReport


@dataclass
class FakeBackend:
    name: str
    mode: str

    def available(self) -> Availability:
        if self.mode == "unavailable":
            return Availability.unavailable(self.name, "not installed")
        return Availability.ok(self.name)

    def recalculate(self, path: Path) -> None:
        if self.mode == "fail":
            path.write_bytes(b"corrupt")
            raise RuntimeError("failed")
        path.write_bytes(path.read_bytes() + b"-recalculated")


def test_chain_uses_pristine_copy_and_never_mutates_input(tmp_path: Path) -> None:
    source = tmp_path / "book.xlsx"
    source.write_bytes(b"original")
    backends = (
        FakeBackend("missing", "unavailable"),
        FakeBackend("broken", "fail"),
        FakeBackend("working", "success"),
    )
    with recalculate(source, backends=backends) as result:
        copy = result.path
        assert result.backend == "working"
        assert copy.read_bytes() == b"original-recalculated"
        assert source.read_bytes() == b"original"
        assert [attempt.backend for attempt in result.attempts] == [
            "missing",
            "broken",
            "working",
        ]
    assert not copy.exists()


def test_all_backend_failures_raise_and_clean(tmp_path: Path) -> None:
    source = tmp_path / "book.xlsx"
    source.write_bytes(b"original")
    with pytest.raises(RecalcFailed) as caught:
        recalculate(
            source,
            backends=(
                FakeBackend("missing", "unavailable"),
                FakeBackend("broken", "fail"),
            ),
        )
    assert len(caught.value.attempts) == 2
    assert source.read_bytes() == b"original"


def test_formulas_backend_rejects_partial_cache_write(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "partial.xlsx"
    book = Workbook()
    book.active["A1"] = 1
    book.active["B1"] = "=A1+1"
    book.save(source)
    monkeypatch.setattr(
        formulas_backend,
        "write_cached_values",
        lambda path, values: CacheReport(2, 1, (), ("CALC!C1",)),
    )
    with pytest.raises(RecalcFailed, match="partial cache write"):
        recalculate(source, backends=(FormulasBackend(),))


def test_formulas_backend_writes_readable_formula_cache(tmp_path: Path) -> None:
    source = tmp_path / "formula.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Calc"
    sheet["A1"] = 6
    sheet["B1"] = 7
    sheet["C1"] = "=A1*B1"
    book.save(source)
    with recalculate(source, backends=(FormulasBackend(),)) as result:
        cached = load_workbook(result.path, data_only=True)
        try:
            assert cached["Calc"]["C1"].value == 42
        finally:
            cached.close()


def test_excel_com_backend_writes_readable_formula_cache(tmp_path: Path) -> None:
    import sys

    from xllib.recalc.backends import ExcelComBackend

    backend = ExcelComBackend()
    availability = backend.available()
    if sys.platform != "win32" or not availability.available:
        pytest.skip(availability.reason or "Excel COM is not available")
    source = tmp_path / "excel-com.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Calc"
    sheet["A1"] = 6
    sheet["B1"] = 7
    sheet["C1"] = "=A1*B1"
    book.save(source)
    with recalculate(source, backends=(backend,)) as result:
        cached = load_workbook(result.path, data_only=True)
        try:
            assert cached["Calc"]["C1"].value == 42
            assert result.backend == "excel_com"
        finally:
            cached.close()
