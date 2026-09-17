"""Machinery fixtures: empty files, hidden sheets, skip/exit 2, waivers."""

from __future__ import annotations

from pathlib import Path

from fixtures.build_fixtures import build_clean_baseline
from openpyxl import Workbook
from openpyxl import load_workbook as openpyxl_load

from xllib.cli import _measure
from xllib.inspect import load_workbook
from xllib.lint.api import lint
from xllib.lint.config import Config, load_config
from xllib.lint.registry import registered_rules
from xllib.lint.rule import Status
from xllib.recalc import recalculate


def test_empty_workbook_does_not_crash(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    Workbook().save(path)
    report = lint(load_workbook(path), Config(), registered_rules())
    assert report.findings
    assert report.exit_code in {0, 2}


def test_hidden_sheet_state_is_reported_on_findings(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "hidden.xlsx")
    book = openpyxl_load(path)
    book["Calc"]["H8"] = "=0.85"
    book["Calc"].sheet_state = "hidden"
    book.create_sheet("Audit").sheet_state = "veryHidden"
    book.save(path)
    with recalculate(path) as result:
        inspected = load_workbook(result.path)
        report = lint(inspected, Config(), registered_rules())
        metrics = _measure(inspected)
    xl002 = next(item for item in report.findings if item.rule_id == "XL002")
    assert xl002.locus.sheet_state == "hidden"
    assert metrics["sheets"] == 4
    assert metrics["sheet_states"]["Calc"] == "hidden"
    assert metrics["sheet_states"]["Audit"] == "veryHidden"


def test_missing_cache_without_recalc_exits_two(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "uncached.xlsx")
    report = lint(load_workbook(path), Config(), registered_rules())
    assert report.skipped >= 1
    assert report.exit_code == 2
    assert all(
        item.status == Status.SKIPPED
        for item in report.findings
        if item.rule_id in {"XL004", "XL005", "XL103"}
    )


def test_named_waiver_suppresses_finding(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "waived.xlsx")
    book = openpyxl_load(path)
    book["Calc"]["H8"] = "=0.85"
    book.save(path)
    config_path = tmp_path / "xllib.toml"
    config_path.write_text(
        """
[[waivers.entries]]
rule = "XL002"
scope = "*"
reason = "fixture"
approver = "owner"
expires = "2099-01-01"
""",
        encoding="utf-8",
    )
    with recalculate(path) as result:
        report = lint(load_workbook(result.path), load_config(config_path), registered_rules())
    xl002 = next(item for item in report.findings if item.rule_id == "XL002")
    assert xl002.status == Status.INFO
    assert xl002.waiver is not None
    assert report.exit_code == 0
