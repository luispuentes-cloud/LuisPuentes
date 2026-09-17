from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook as OpenpyxlWorkbook

from xllib.inspect import (
    detect_period_axis,
    estimate_section_count,
    infer_formula_runs,
    load_workbook,
)


def test_detects_period_axis_formula_runs_and_sections(tmp_path: Path) -> None:
    path = tmp_path / "shapes.xlsx"
    book = OpenpyxlWorkbook()
    sheet = book.active
    sheet.title = "Calc"
    sheet.append(["Model"])
    sheet.append(["Period", None, None, "FY26", "FY27", "FY28"])
    sheet["C4"] = 2
    sheet["D4"] = 10
    sheet["E4"] = 20
    sheet["F4"] = 30
    sheet["D5"] = "=D4*$C$4"
    sheet["E5"] = "=E4*$C$4"
    sheet["F5"] = "=F4*$C$4"
    sheet["G5"] = "=SUM(D5:F5)"
    sheet["A7"] = "Notes"
    book.save(path)

    inspected = load_workbook(path)
    calc = inspected.sheet("Calc")
    assert calc is not None

    axis = detect_period_axis(calc)
    assert axis is not None
    assert (axis.row, axis.start_column, axis.end_column) == (2, 4, 6)
    assert axis.labels == ("FY26", "FY27", "FY28")

    runs = infer_formula_runs(calc)
    assert len(runs) == 1
    assert (runs[0].start_column, runs[0].end_column) == (4, 7)
    assert len(set(runs[0].shapes[:3])) == 1
    assert runs[0].shapes[3] != runs[0].shapes[0]
    assert estimate_section_count(calc) == 3


def test_hidden_sheet_is_available_to_all_inspection_heuristics(
    tmp_path: Path,
) -> None:
    path = tmp_path / "hidden.xlsx"
    book = OpenpyxlWorkbook()
    hidden = book.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "FY26"
    hidden["B1"] = "FY27"
    hidden["A2"] = "=A3"
    hidden["B2"] = "=B3"
    book.save(path)

    inspected = load_workbook(path)
    sheet = inspected.sheet("Hidden")
    assert sheet is not None

    assert sheet.state == "hidden"
    assert detect_period_axis(sheet) is not None
    assert len(infer_formula_runs(sheet)) == 1
