"""Generate Phase 0 fixtures at test time. No committed binaries."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

_INPUT_FONT = Font(color="0000FF")
_INPUT_FILL = PatternFill(fill_type="solid", fgColor="DCE6F1")


def _style_input(cell: object) -> None:
    cell.font = _INPUT_FONT  # type: ignore[attr-defined]
    cell.fill = _INPUT_FILL  # type: ignore[attr-defined]


def build_clean_baseline(path: Path) -> Path:
    """A three-sheet workbook that should produce no VIOLATION findings."""
    book = Workbook()
    inputs = book.active
    inputs.title = "Inputs"
    inputs["B4"] = "Hours factor"
    inputs["C4"] = 0.85
    _style_input(inputs["C4"])
    inputs["B5"] = "Legend"
    inputs["C5"] = "input / calculation / check"

    calc = book.create_sheet("Calc")
    calc["C3"] = "Metric"
    calc["D3"] = "FY26"
    calc["E3"] = "FY27"
    calc["F3"] = "FY28"
    calc["C4"] = "Hours"
    calc["D4"] = 10
    calc["E4"] = 10
    calc["F4"] = 10
    _style_input(calc["D4"])
    _style_input(calc["E4"])
    _style_input(calc["F4"])
    calc["C5"] = "Result"
    calc["D5"] = "=D4*HoursFactor"
    calc["E5"] = "=E4*HoursFactor"
    calc["F5"] = "=F4*HoursFactor"
    calc["C6"] = "Check"
    calc["D6"] = "=D5-D5"

    summary = book.create_sheet("Summary")
    summary["C3"] = "Metric"
    summary["D3"] = "FY26"
    summary["E3"] = "FY27"
    summary["F3"] = "FY28"
    summary["D4"] = "=Calc!D5"
    summary["E4"] = "=Calc!E5"
    summary["F4"] = "=Calc!F5"

    book.defined_names.add(DefinedName("HoursFactor", attr_text="'Inputs'!$C$4"))
    book.defined_names.add(DefinedName("chk_total_tie", attr_text="'Calc'!$D$6"))
    book.save(path)
    return path
