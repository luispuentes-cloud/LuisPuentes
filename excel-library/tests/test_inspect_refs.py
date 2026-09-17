from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl.workbook.defined_name import DefinedName

from xllib.inspect import build_reference_graph, load_workbook, walk_formula


def test_reference_graph_resolves_names_ranges_and_external_edges(
    tmp_path: Path,
) -> None:
    path = tmp_path / "references.xlsx"
    book = OpenpyxlWorkbook()
    inputs = book.active
    inputs.title = "Inputs"
    inputs["B2"] = 10
    inputs["B3"] = 20
    inputs["B4"] = 30
    calc = book.create_sheet("Calc")
    calc["D7"] = "=SUM(Inputs!B2:B4)+DRIVER+'[other.xlsx]Sheet1'!A1"
    book.defined_names.add(DefinedName("DRIVER", attr_text="'Inputs'!$B$2"))
    book.save(path)

    inspected = load_workbook(path)
    graph = build_reference_graph(inspected)

    assert len(graph.edges) == 3
    assert graph.edges[0].target.text == "Inputs!B2:B4"
    named_edge = next(edge for edge in graph.edges if edge.via_defined_name)
    assert named_edge.via_defined_name == "DRIVER"
    assert named_edge.target.sheet == "Inputs"
    assert named_edge.target.is_single_cell
    assert len(graph.aggregate_ranges) == 1
    assert graph.aggregate_ranges[0].function == "SUM"
    assert graph.aggregate_ranges[0].target.min_row == 2
    assert graph.aggregate_ranges[0].target.max_row == 4
    assert len(graph.external_edges) == 1


def test_token_walk_marks_resolved_defined_name(tmp_path: Path) -> None:
    path = tmp_path / "name.xlsx"
    book = OpenpyxlWorkbook()
    book.active["A1"] = 4
    book.active["A2"] = "=RATE"
    book.defined_names.add(DefinedName("RATE", attr_text="'Sheet'!$A$1"))
    book.save(path)

    inspected = load_workbook(path)
    walk = walk_formula("=rate", workbook=inspected, sheet_index=0)

    assert walk.range_tokens[0].defined_name == "RATE"
