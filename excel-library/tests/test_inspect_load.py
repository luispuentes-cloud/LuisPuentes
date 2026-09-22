from __future__ import annotations

import time
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook as openpyxl_load_workbook
from openpyxl.workbook.defined_name import DefinedName

from xllib.inspect import (
    Capability,
    capabilities,
    detect_period_axis,
    estimate_section_count,
    infer_formula_runs,
    load_workbook,
)

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _set_formula_cache(path: Path, coordinate: str, value: str) -> None:
    with ZipFile(path, "r") as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(members["xl/worksheets/sheet1.xml"])
    cell = root.find(f".//{{{_MAIN_NS}}}c[@r='{coordinate}']")
    assert cell is not None
    cached = cell.find(f"{{{_MAIN_NS}}}v")
    if cached is None:
        cached = ET.SubElement(cell, f"{{{_MAIN_NS}}}v")
    cached.text = value
    members["xl/worksheets/sheet1.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)


def test_dual_load_pairs_formula_and_cached_value(tmp_path: Path) -> None:
    path = tmp_path / "cached.xlsx"
    book = OpenpyxlWorkbook()
    sheet = book.active
    sheet.title = "Calc"
    sheet["A1"] = 6
    sheet["B1"] = "=A1*7"
    book.save(path)
    _set_formula_cache(path, "B1", "42")

    inspected = load_workbook(path)
    calc = inspected.sheet("calc")
    assert calc is not None
    formula_cell = calc.cell(1, 2)

    assert formula_cell is not None
    assert formula_cell.formula == "=A1*7"
    assert formula_cell.cached_value == 42
    assert Capability.CACHED_VALUES in capabilities(inspected)


def test_load_preserves_hidden_states_and_defined_names(tmp_path: Path) -> None:
    path = tmp_path / "states.xlsx"
    book = OpenpyxlWorkbook()
    inputs = book.active
    inputs.title = "Inputs"
    inputs["B2"] = 0.85
    hidden = book.create_sheet("HiddenCalc")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "=DRIVER"
    very_hidden = book.create_sheet("Audit")
    very_hidden.sheet_state = "veryHidden"
    very_hidden["A1"] = "retained"
    book.defined_names.add(DefinedName("DRIVER", attr_text="'Inputs'!$B$2"))
    book.save(path)

    inspected = load_workbook(path)

    hidden_calc = inspected.sheet("HiddenCalc")
    assert hidden_calc is not None
    assert tuple(sheet.state for sheet in inspected.sheets) == (
        "visible",
        "hidden",
        "veryHidden",
    )
    hidden_formula = hidden_calc.cell(1, 1)
    assert hidden_formula is not None
    assert hidden_formula.formula == "=DRIVER"
    driver = inspected.resolve_name("driver", sheet_index=1)
    assert driver is not None
    assert driver.references[0].sheet == "Inputs"
    assert driver.references[0].is_single_cell


def test_missing_formula_cache_is_not_claimed(tmp_path: Path) -> None:
    path = tmp_path / "uncalculated.xlsx"
    book = OpenpyxlWorkbook()
    book.active["A1"] = "=1+1"
    book.save(path)

    inspected = load_workbook(path)

    assert Capability.CACHED_VALUES not in capabilities(inspected)


def test_a_far_corner_cell_does_not_cost_the_whole_rectangle(tmp_path: Path) -> None:
    """Two real cells used to cost ten million coordinates, twice over.

    `iter_rows()` walks the declared rectangle and creates a cell object for
    every coordinate inside it, and the span heuristics then walked the same
    rectangle again looking for content. A single stray value in the far
    corner of a sheet is a common file shape, and Phase 0 promises to run on
    an arbitrary `.xlsx`.

    The bound is deliberately loose. What is asserted is the difference
    between a cost set by the content and one set by the declared dimension,
    not a benchmark. Measured on this fixture: 364s before the fix, and under
    a second after it.
    """
    path = tmp_path / "far-corner.xlsx"
    book = OpenpyxlWorkbook()
    sheet = book.active
    sheet.title = "Sparse"
    sheet["A1"] = 1
    sheet.cell(row=20000, column=500, value=2)
    book.save(path)

    started = time.perf_counter()
    sparse = load_workbook(path).sheet("Sparse")
    assert sparse is not None
    detect_period_axis(sparse)
    infer_formula_runs(sparse)
    sections = estimate_section_count(sparse)
    elapsed = time.perf_counter() - started

    assert len(sparse.cells) == 2
    assert sparse.populated_rows == (1, 20000)
    assert sections == 2
    assert elapsed < 10


def test_the_loader_reads_the_same_cells_a_rectangle_walk_would(tmp_path: Path) -> None:
    """`_stored_cells` reads a private openpyxl attribute, so pin it to the API.

    The control for the test above: making the read cheap is only a fix if it
    also reads the same thing.
    """
    path = tmp_path / "dense.xlsx"
    book = OpenpyxlWorkbook()
    sheet = book.active
    sheet.title = "Dense"
    sheet["A1"] = "label"
    sheet["B1"] = 2
    sheet["D2"] = "=B1*2"
    sheet["A4"] = 0
    book.save(path)

    loaded = load_workbook(path).sheet("Dense")
    assert loaded is not None
    walked = openpyxl_load_workbook(path)["Dense"]

    assert {cell.coordinate for cell in loaded.cells} == {
        cell.coordinate for row in walked.iter_rows() for cell in row if cell.value is not None
    }
