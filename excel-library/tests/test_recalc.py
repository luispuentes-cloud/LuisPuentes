"""Tests for xllib.recalc.

The chain is proved with fake backends: ordering, fallback, the temp copy, the
input staying untouched, and cleanup are all properties of `recalculate`
itself, and testing them through `formulas` or Excel would make them
untestable in CI. The real backends are covered by availability smoke tests
plus, for Excel COM, a fake Application that proves the mandatory flags are
set.
"""

from __future__ import annotations

import hashlib
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from xllib.recalc import (
    Availability,
    ExcelComBackend,
    FormulasBackend,
    LibreOfficeBackend,
    Outcome,
    RecalcBackend,
    RecalcFailed,
    default_backends,
    recalculate,
    write_cached_values,
)
from xllib.recalc.backend import BackendError

# --------------------------------------------------------------------------
# fakes and fixtures
# --------------------------------------------------------------------------


@dataclass
class FakeBackend:
    """A backend whose behaviour is declared rather than computed.

    `reason` makes it unavailable, `error` makes it fail, and `writes` is the
    content it leaves in the copy it was handed.
    """

    name: str
    reason: str | None = None
    error: Exception | None = None
    writes: bytes | None = None
    calls: list[Path] = field(default_factory=list)
    received: list[bytes] = field(default_factory=list)

    def available(self) -> Availability:
        if self.reason is not None:
            return Availability.unavailable(self.name, self.reason)
        return Availability.ok(self.name)

    def recalculate(self, path: Path) -> None:
        self.calls.append(path)
        self.received.append(path.read_bytes())
        if self.writes is not None:
            path.write_bytes(self.writes)
        if self.error is not None:
            raise self.error


ORIGINAL = b"the original workbook bytes"


@pytest.fixture
def source(tmp_path: Path) -> Path:
    target = tmp_path / "model.xlsx"
    target.write_bytes(ORIGINAL)
    return target


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workspace_of(result_path: Path) -> Path:
    """The temp workspace root, two levels above the staged copy."""
    return result_path.parent.parent


# --------------------------------------------------------------------------
# ordering and fallback
# --------------------------------------------------------------------------


def test_first_available_backend_wins_and_the_rest_are_never_consulted(
    source: Path,
) -> None:
    first = FakeBackend("first")
    second = FakeBackend("second")

    with recalculate(source, backends=[first, second]) as result:
        assert result.backend == "first"

    assert len(first.calls) == 1
    assert second.calls == []


def test_an_unavailable_backend_is_skipped_with_its_reason_recorded(
    source: Path,
) -> None:
    unavailable = FakeBackend("unavailable", reason="not installed here")
    fallback = FakeBackend("fallback")

    with recalculate(source, backends=[unavailable, fallback]) as result:
        assert result.backend == "fallback"

    first, second = result.attempts
    assert (first.backend, first.outcome, first.detail) == (
        "unavailable",
        Outcome.UNAVAILABLE,
        "not installed here",
    )
    assert (second.backend, second.outcome) == ("fallback", Outcome.SUCCEEDED)
    assert unavailable.calls == []


def test_a_failing_backend_falls_through_to_the_next_one(source: Path) -> None:
    failing = FakeBackend("failing", error=RuntimeError("engine blew up"))
    fallback = FakeBackend("fallback")

    with recalculate(source, backends=[failing, fallback]) as result:
        assert result.backend == "fallback"

    attempt = result.attempts[0]
    assert attempt.outcome is Outcome.FAILED
    assert attempt.detail == "engine blew up"
    assert isinstance(attempt.error, RuntimeError)


def test_every_backend_is_tried_before_the_chain_gives_up(source: Path) -> None:
    backends = [
        FakeBackend("a", reason="absent"),
        FakeBackend("b", error=BackendError("b", "crashed")),
        FakeBackend("c", error=RuntimeError("also crashed")),
    ]

    with pytest.raises(RecalcFailed) as raised:
        recalculate(source, backends=backends)

    assert [a.backend for a in raised.value.attempts] == ["a", "b", "c"]
    assert [a.outcome for a in raised.value.attempts] == [
        Outcome.UNAVAILABLE,
        Outcome.FAILED,
        Outcome.FAILED,
    ]
    message = str(raised.value)
    assert "a unavailable (absent)" in message
    assert "b failed" in message


def test_an_empty_chain_fails_rather_than_returning_an_unrecalculated_copy(
    source: Path,
) -> None:
    with pytest.raises(RecalcFailed) as raised:
        recalculate(source, backends=[])

    assert raised.value.attempts == ()


def test_the_default_chain_is_formulas_then_libreoffice_then_excel() -> None:
    assert [backend.name for backend in default_backends()] == [
        "formulas",
        "libreoffice",
        "excel_com",
    ]


# --------------------------------------------------------------------------
# the temp copy
# --------------------------------------------------------------------------


def test_the_backend_works_on_a_copy_that_keeps_the_original_name(
    source: Path,
) -> None:
    backend = FakeBackend("copy")

    with recalculate(source, backends=[backend]) as result:
        assert result.path != source
        assert source.parent not in result.path.parents
        assert result.path.name == source.name
        assert backend.calls == [result.path]
        assert backend.received == [ORIGINAL]


def test_the_input_is_untouched_even_when_a_backend_rewrites_its_copy(
    source: Path,
) -> None:
    before = digest(source)
    backend = FakeBackend("rewrites", writes=b"recalculated")

    with recalculate(source, backends=[backend]) as result:
        assert result.path.read_bytes() == b"recalculated"
        assert digest(source) == before

    assert digest(source) == before
    assert source.read_bytes() == ORIGINAL


def test_each_attempt_starts_from_a_pristine_copy(source: Path) -> None:
    corrupting = FakeBackend(
        "corrupting", writes=b"half-written garbage", error=RuntimeError("died midway")
    )
    fallback = FakeBackend("fallback")

    with recalculate(source, backends=[corrupting, fallback]) as result:
        assert result.backend == "fallback"

    assert fallback.received == [ORIGINAL]
    assert corrupting.calls[0] != fallback.calls[0]


def test_result_reports_the_source_and_the_winning_backend(source: Path) -> None:
    with recalculate(source, backends=[FakeBackend("winner")]) as result:
        assert result.source == source.resolve()
        assert result.backend == "winner"
        assert [a.outcome for a in result.attempts] == [Outcome.SUCCEEDED]


def test_a_hostile_backend_name_cannot_escape_the_workspace(source: Path) -> None:
    backend = FakeBackend("../../escape")

    with recalculate(source, backends=[backend]) as result:
        assert result.path.parent.name == "00-escape"
        assert result.path.is_relative_to(workspace_of(result.path))


# --------------------------------------------------------------------------
# cleanup
# --------------------------------------------------------------------------


def test_leaving_the_context_removes_the_workspace(source: Path) -> None:
    with recalculate(source, backends=[FakeBackend("cleanup")]) as result:
        copy = result.path
        workspace = workspace_of(copy)
        assert copy.exists()

    assert not copy.exists()
    assert not workspace.exists()


def test_close_is_idempotent_and_the_path_refuses_to_answer_afterwards(
    source: Path,
) -> None:
    result = recalculate(source, backends=[FakeBackend("cleanup")])
    workspace = workspace_of(result.path)

    result.close()
    result.close()

    assert result.closed
    assert not workspace.exists()
    with pytest.raises(RuntimeError, match="already been discarded"):
        _ = result.path
    # Diagnostics survive the copy.
    assert result.backend == "cleanup"
    assert [a.outcome for a in result.attempts] == [Outcome.SUCCEEDED]


def test_the_workspace_is_removed_when_the_whole_chain_fails(source: Path) -> None:
    staged: list[Path] = []

    @dataclass
    class Recording(FakeBackend):
        def recalculate(self, path: Path) -> None:
            staged.append(path)
            raise RuntimeError("no good")

    with pytest.raises(RecalcFailed):
        recalculate(source, backends=[Recording("recording")])

    assert staged, "the backend should have been handed a copy"
    assert not workspace_of(staged[0]).exists()


def test_the_workspace_is_removed_when_a_backend_raises_through_the_chain(
    source: Path,
) -> None:
    staged: list[Path] = []

    class Interrupting:
        name = "interrupting"

        def available(self) -> Availability:
            return Availability.ok(self.name)

        def recalculate(self, path: Path) -> None:
            staged.append(path)
            raise KeyboardInterrupt

    with pytest.raises(KeyboardInterrupt):
        recalculate(source, backends=[Interrupting()])

    assert not workspace_of(staged[0]).exists()


def test_a_missing_input_is_rejected_before_any_workspace_is_made(
    tmp_path: Path,
) -> None:
    with pytest.raises(FileNotFoundError):
        recalculate(tmp_path / "absent.xlsx", backends=[FakeBackend("never")])


def test_a_directory_is_not_a_workbook(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="not a file"):
        recalculate(tmp_path, backends=[FakeBackend("never")])


# --------------------------------------------------------------------------
# the protocol
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "backend",
    [FormulasBackend(), LibreOfficeBackend(), ExcelComBackend(), FakeBackend("fake")],
    ids=["formulas", "libreoffice", "excel_com", "fake"],
)
def test_backends_satisfy_the_protocol(backend: Any) -> None:
    assert isinstance(backend, RecalcBackend)


@pytest.mark.parametrize(
    "backend",
    [FormulasBackend(), LibreOfficeBackend(), ExcelComBackend()],
    ids=["formulas", "libreoffice", "excel_com"],
)
def test_availability_is_answered_without_raising_and_explains_itself(
    backend: Any,
) -> None:
    availability = backend.available()

    assert availability.backend == backend.name
    assert isinstance(availability.available, bool)
    # A reason is present exactly when the backend cannot run.
    assert availability.available is not bool(availability.reason)


# --------------------------------------------------------------------------
# formulas backend
# --------------------------------------------------------------------------


def test_formulas_reports_itself_unavailable_when_the_package_is_missing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "xllib.recalc.backends.formulas.importlib.util.find_spec", lambda name: None
    )
    availability = FormulasBackend().available()

    assert not availability.available
    assert "formulas package is not installed" in availability.reason


def test_formulas_progress_bar_never_reaches_stderr(
    capsys: pytest.CaptureFixture[str],
) -> None:
    from xllib.recalc.backends.formulas import _quiet_progress

    with _quiet_progress():
        print("100%|##########| 12/12", file=sys.stderr)

    assert capsys.readouterr().err == ""


def test_formulas_solution_keys_are_mapped_case_insensitively_by_sheet() -> None:
    from xllib.recalc.backends.formulas import _values_by_sheet

    values = _values_by_sheet(
        {
            "'[MODEL.XLSX]CALC'!D7": 42.0,
            "'[MODEL.XLSX]Calc'!D8": "ok",
            "'[OTHER.XLSX]CALC'!D9": 1.0,
            "'[MODEL.XLSX]CALC'!D10:E10": [[1.0, 2.0]],
            "'[MODEL.XLSX]'!FTE_TARGET": 5.0,
        },
        "model.xlsx",
    )

    assert values == {"CALC": {"D7": 42.0, "D8": "ok"}}


def test_formulas_scalars_are_unwrapped_and_arrays_are_left_alone() -> None:
    from xllib.recalc.backends.formulas import _UNSET, _scalar

    class Range:
        def __init__(self, value: Any) -> None:
            self.value = value

    class Array:
        def __init__(self, rows: Any) -> None:
            self._rows = rows

        def tolist(self) -> Any:
            return self._rows

    assert _scalar(Range(Array([[6.0]]))) == 6.0
    assert _scalar(Range(Array([["ok"]]))) == "ok"
    assert _scalar(Array([[1.0, 2.0]])) is _UNSET
    assert _scalar(Array([])) is _UNSET
    assert _scalar(True) is True
    assert _scalar(None) is None
    assert _scalar(object()) is _UNSET


def test_formulas_numpy_scalars_are_demoted_to_python_types() -> None:
    from xllib.recalc.backends.formulas import _scalar

    numpy = pytest.importorskip("numpy", reason="numpy ships with formulas")

    assert _scalar(numpy.float64(6.0)) == 6.0
    assert type(_scalar(numpy.float64(6.0))) is float
    assert _scalar(numpy.str_("ok")) == "ok"
    assert type(_scalar(numpy.bool_(True))) is bool
    assert _scalar(numpy.array([[6.0]])) == 6.0


@pytest.mark.parametrize("value", ["#DIV/0!", "#VALUE!"])
def test_formulas_error_tokens_survive_as_errors(value: str) -> None:
    from xllib.recalc.backends.formulas import _scalar

    class Token:
        def __init__(self, text: str) -> None:
            self._text = text

        def __str__(self) -> str:
            return self._text

    assert _scalar(Token(value)) == value


# --------------------------------------------------------------------------
# libreoffice backend
# --------------------------------------------------------------------------


def test_libreoffice_says_where_it_looked_when_soffice_is_absent(
    tmp_path: Path,
) -> None:
    availability = LibreOfficeBackend(soffice=tmp_path / "soffice").available()

    assert not availability.available
    assert "not found on PATH" in availability.reason


def test_libreoffice_refuses_to_run_without_soffice_rather_than_half_trying(
    tmp_path: Path, source: Path
) -> None:
    backend = LibreOfficeBackend(soffice=tmp_path / "soffice")

    with pytest.raises(BackendError, match="not found on PATH"):
        backend.recalculate(source)

    assert source.read_bytes() == ORIGINAL


def test_libreoffice_reports_a_missing_uno_bridge_separately_from_a_missing_install(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    pretend_soffice = tmp_path / "soffice"
    pretend_soffice.write_text("", encoding="utf-8")
    monkeypatch.setattr(
        "xllib.recalc.backends.libreoffice.importlib.util.find_spec",
        lambda name: None,
    )

    availability = LibreOfficeBackend(soffice=pretend_soffice).available()

    assert not availability.available
    assert "python-uno bridge is not importable" in availability.reason


# --------------------------------------------------------------------------
# excel com backend
# --------------------------------------------------------------------------


class FakeWorkbook:
    def __init__(self, application: FakeExcel) -> None:
        self._application = application
        self.saves = 0
        self.closed_with: list[bool] = []

    def Save(self) -> None:
        if self._application.save_error is not None:
            raise self._application.save_error
        self.saves += 1

    def Close(self, save_changes: bool) -> None:
        self.closed_with.append(save_changes)


class FakeWorkbooks:
    def __init__(self, application: FakeExcel) -> None:
        self._application = application

    def Open(self, filename: str, update_links: int, read_only: bool) -> FakeWorkbook:
        application = self._application
        application.opened = (filename, update_links, read_only)
        # Captured at open time: a window must never have had the chance to show.
        application.flags_at_open = (
            application.Visible,
            application.DisplayAlerts,
            application.EnableEvents,
        )
        application.book = FakeWorkbook(application)
        return application.book


class FakeExcel:
    """Stands in for Excel.Application, with Excel's real defaults."""

    def __init__(
        self, *, calculate_error: Exception | None = None, save_error: Exception | None = None
    ) -> None:
        self.Visible = True
        self.DisplayAlerts = True
        self.EnableEvents = True
        self.Workbooks = FakeWorkbooks(self)
        self.calculate_error = calculate_error
        self.save_error = save_error
        self.opened: tuple[str, int, bool] | None = None
        self.flags_at_open: tuple[bool, bool, bool] | None = None
        self.book: FakeWorkbook | None = None
        self.rebuilds = 0
        self.quits = 0
        self.released = 0

    def CalculateFullRebuild(self) -> None:
        if self.calculate_error is not None:
            raise self.calculate_error
        self.rebuilds += 1

    def Quit(self) -> None:
        self.quits += 1

    def release(self) -> None:
        self.released += 1


def excel_backend(application: FakeExcel) -> ExcelComBackend:
    return ExcelComBackend(dispatch=lambda: (application, application.release))


def test_excel_is_hidden_and_silent_before_the_workbook_is_opened(
    source: Path,
) -> None:
    application = FakeExcel()

    excel_backend(application).recalculate(source)

    assert application.flags_at_open == (False, False, False)
    assert application.opened == (str(source), 0, False)


def test_excel_recalculates_saves_closes_and_quits(source: Path) -> None:
    application = FakeExcel()

    excel_backend(application).recalculate(source)

    assert application.rebuilds == 1
    assert application.book is not None
    assert application.book.saves == 1
    assert application.book.closed_with == [False]
    assert (application.quits, application.released) == (1, 1)


@pytest.mark.parametrize(
    "failure",
    [
        {"calculate_error": RuntimeError("calc failed")},
        {"save_error": RuntimeError("save failed")},
    ],
    ids=["calculate", "save"],
)
def test_excel_is_closed_and_quit_even_when_recalculation_fails(
    source: Path, failure: dict[str, Exception]
) -> None:
    application = FakeExcel(**failure)

    with pytest.raises(RuntimeError):
        excel_backend(application).recalculate(source)

    assert application.book is not None
    assert application.book.closed_with == [False]
    assert (application.quits, application.released) == (1, 1)


def test_a_failing_excel_quit_still_releases_the_com_apartment(source: Path) -> None:
    application = FakeExcel()
    application.Quit = lambda: (_ for _ in ()).throw(RuntimeError("quit failed"))  # type: ignore[method-assign]

    with pytest.raises(RuntimeError, match="quit failed"):
        excel_backend(application).recalculate(source)

    assert application.released == 1


@pytest.mark.skipif(sys.platform == "win32", reason="Excel COM is available here")
def test_excel_is_unavailable_off_windows() -> None:
    availability = ExcelComBackend().available()

    assert not availability.available
    assert "requires Windows" in availability.reason


# --------------------------------------------------------------------------
# cached value injection
# --------------------------------------------------------------------------


@pytest.fixture
def formula_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "fixture.xlsx"
    workbook = Workbook()
    calc = workbook.active
    calc.title = "Calc"
    calc["A1"], calc["A2"], calc["A3"] = 1, 2, 3
    calc["B4"] = "=SUM(A1:A3)"
    calc["B4"].number_format = "#,##0.00"
    calc["B5"] = '=IF(B4=6,"ok","bad")'
    calc["B6"] = "=B4>0"
    calc["B7"] = "=1/0"
    calc["B8"] = "not a formula"
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "prose"
    workbook.save(path)
    return path


def test_injected_values_are_readable_and_typed(formula_workbook: Path) -> None:
    report = write_cached_values(
        formula_workbook,
        {"CALC": {"B4": 6.0, "B5": "ok", "B6": True, "B7": "#DIV/0!"}},
    )

    assert (report.formula_cells, report.written) == (4, 4)

    values = load_workbook(formula_workbook, data_only=True)["Calc"]
    assert values["B4"].value == 6.0
    assert values["B5"].value == "ok"
    assert values["B6"].value is True
    assert values["B7"].value == "#DIV/0!"


def test_injection_keeps_the_formulas_and_the_number_formats(
    formula_workbook: Path,
) -> None:
    write_cached_values(formula_workbook, {"Calc": {"B4": 6.0}})

    calc = load_workbook(formula_workbook)["Calc"]
    assert calc["B4"].value == "=SUM(A1:A3)"
    assert calc["B4"].number_format == "#,##0.00"
    assert calc["B8"].value == "not a formula"
    assert load_workbook(formula_workbook)["Notes"]["A1"].value == "prose"


def test_sheet_and_reference_keys_are_matched_case_insensitively(
    formula_workbook: Path,
) -> None:
    # This is the shape `formulas` hands back: the sheet name upper-cased.
    report = write_cached_values(formula_workbook, {"CALC": {"b4": 6.0}})

    assert report.written == 1
    assert load_workbook(formula_workbook, data_only=True)["Calc"]["B4"].value == 6.0


def test_a_sheet_that_is_not_in_the_workbook_is_reported_not_guessed(
    formula_workbook: Path,
) -> None:
    report = write_cached_values(formula_workbook, {"Ghost": {"A1": 1.0}})

    assert report.unmatched_sheets == ("GHOST",)
    assert report.written == 0


def test_values_for_cells_without_formulas_are_ignored(
    formula_workbook: Path,
) -> None:
    report = write_cached_values(formula_workbook, {"Calc": {"B8": 99.0, "Z99": 1.0}})

    assert report.written == 0
    assert load_workbook(formula_workbook)["Calc"]["B8"].value == "not a formula"


def test_reinjection_replaces_the_previous_value(formula_workbook: Path) -> None:
    write_cached_values(formula_workbook, {"Calc": {"B4": 6.0}})
    write_cached_values(formula_workbook, {"Calc": {"B4": 7.0}})

    assert load_workbook(formula_workbook, data_only=True)["Calc"]["B4"].value == 7.0


def test_a_number_whose_repr_is_not_a_number_is_still_written_as_one(
    formula_workbook: Path,
) -> None:
    class Numpyish(float):
        def __repr__(self) -> str:
            return "np.float64(6.0)"

    report = write_cached_values(formula_workbook, {"Calc": {"B4": Numpyish(6.0)}})

    assert report.written == 1
    # The assertion that matters is that the workbook is still readable at all.
    assert load_workbook(formula_workbook, data_only=True)["Calc"]["B4"].value == 6.0


def test_unrepresentable_values_are_left_unwritten(formula_workbook: Path) -> None:
    report = write_cached_values(formula_workbook, {"Calc": {"B4": None, "B5": object()}})

    assert report.written == 1
    assert "CALC!B5" in report.unwritten


def test_partial_injection_lists_unwritten_formula_cells(formula_workbook: Path) -> None:
    report = write_cached_values(formula_workbook, {"Calc": {"B4": 6.0}})

    assert report.formula_cells == 4
    assert report.written == 1
    assert len(report.unwritten) == 3


# --------------------------------------------------------------------------
# platform integration, skipped unless the engine is installed
# --------------------------------------------------------------------------


def test_formulas_backend_produces_cached_values_end_to_end(
    formula_workbook: Path,
) -> None:
    pytest.importorskip("formulas", reason="the formulas backend is not installed")
    before = digest(formula_workbook)

    with recalculate(formula_workbook, backends=[FormulasBackend()]) as result:
        assert result.backend == "formulas"
        recalculated = load_workbook(result.path, data_only=True)["Calc"]
        assert recalculated["B4"].value == 6
        assert load_workbook(result.path)["Calc"]["B4"].number_format == "#,##0.00"

    assert digest(formula_workbook) == before
