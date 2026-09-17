"""Full-registry finding sets for the clean baseline and isolated mutations."""

from __future__ import annotations

from pathlib import Path

from fixtures.build_fixtures import build_clean_baseline
from openpyxl import load_workbook as openpyxl_load
from openpyxl.styles import Font, PatternFill

from xllib.inspect import load_workbook
from xllib.lint.api import lint
from xllib.lint.config import Config
from xllib.lint.registry import registered_rules
from xllib.lint.rule import Status
from xllib.recalc import recalculate

_INPUT_FONT = Font(color="0000FF")
_INPUT_FILL = PatternFill(fill_type="solid", fgColor="DCE6F1")


def _lint_recalculated(path: Path, config: Config | None = None):
    with recalculate(path) as result:
        return lint(load_workbook(result.path), config or Config(), registered_rules())


def _violation_ids(report) -> list[str]:
    return [
        finding.rule_id for finding in report.findings if finding.status == Status.VIOLATION
    ]


def test_each_registered_rule_has_a_doc_page() -> None:
    root = Path(__file__).resolve().parents[1]
    for rule in registered_rules():
        assert (root / rule.doc).is_file(), rule.doc


def test_clean_baseline_full_registry_has_no_violations(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "clean.xlsx")
    report = _lint_recalculated(path)
    assert _violation_ids(report) == []
    assert report.skipped == 0
    assert report.exit_code == 0


def test_xl001_mutation_finding_set_is_exact(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "xl001.xlsx")
    book = openpyxl_load(path)
    book["Inputs"]["B4"] = None
    book.save(path)
    report = _lint_recalculated(path)
    assert _violation_ids(report) == ["XL001"]


def test_xl002_mutation_finding_set_is_exact(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "xl002.xlsx")
    book = openpyxl_load(path)
    book["Calc"]["H8"] = "=0.85"
    book.save(path)
    report = _lint_recalculated(path)
    assert _violation_ids(report) == ["XL002"]
    assert report.findings[0].evidence["literal"] == "0.85"


def test_xl003_mutation_finding_set_is_exact(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "xl003.xlsx")
    book = openpyxl_load(path)
    book["Calc"]["F5"] = "=SUM(D5:E5)"
    book.save(path)
    report = _lint_recalculated(path)
    assert _violation_ids(report) == ["XL003"]


def test_xl004_mutation_finding_set_is_exact(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "xl004.xlsx")
    book = openpyxl_load(path)
    book["Calc"]["H12"] = "N/A"
    book["Calc"]["H12"].number_format = "$#,##0"
    book.save(path)
    report = _lint_recalculated(path)
    assert _violation_ids(report) == ["XL004"]


def test_xl005_mutation_finding_set_is_exact(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "xl005.xlsx")
    book = openpyxl_load(path)
    cell = book["Calc"]["H14"]
    cell.value = "n/a"
    cell.font = _INPUT_FONT
    cell.fill = _INPUT_FILL
    book["Summary"]["G8"] = "=SUM(Calc!H14:H14)"
    book.save(path)
    report = _lint_recalculated(path)
    assert _violation_ids(report) == ["XL005"]


def test_xl006_mutation_finding_set_is_exact(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "xl006.xlsx")
    book = openpyxl_load(path)
    summary = book["Summary"]
    for column, label in enumerate(("FY26", "FY27", "FY28"), start=5):
        summary.cell(3, column, label)
    summary["D3"] = None
    book.save(path)
    report = _lint_recalculated(path)
    assert _violation_ids(report) == ["XL006"]


def test_xl101_mutation_finding_set_is_exact(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "xl101.xlsx")
    book = openpyxl_load(path)
    for index, row in enumerate((7, 9, 11, 13, 15, 17, 19)):
        book["Inputs"].cell(row, 1, f"Block {index}")
    book.save(path)
    report = _lint_recalculated(path)
    assert _violation_ids(report) == ["XL101"]


def test_xl102_mutation_finding_set_is_exact(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "xl102.xlsx")
    book = openpyxl_load(path)
    del book.defined_names["chk_total_tie"]
    book.save(path)
    report = _lint_recalculated(path)
    assert _violation_ids(report) == ["XL102"]


def test_xl103_mutation_finding_set_is_exact(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "xl103.xlsx")
    book = openpyxl_load(path)
    book["Calc"]["D6"] = "=1"
    book.save(path)
    report = _lint_recalculated(path)
    assert _violation_ids(report) == ["XL103"]


def test_xl104_mutation_finding_set_is_exact(tmp_path: Path) -> None:
    path = build_clean_baseline(tmp_path / "xl104.xlsx")
    book = openpyxl_load(path)
    cell = book["Inputs"]["C4"]
    cell.font = Font()
    cell.fill = PatternFill()
    book.save(path)
    report = _lint_recalculated(path)
    assert _violation_ids(report) == ["XL104"]
