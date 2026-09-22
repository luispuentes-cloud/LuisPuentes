"""Command-line interface."""

from __future__ import annotations

import argparse
import json
import sys
from collections.abc import Sequence
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl.utils.exceptions import InvalidFileException

from xllib.inspect import (
    assert_readable,
    estimate_section_count,
    load_workbook,
    normalize_formula_shape,
    walk_formula,
)
from xllib.lint import discover_config, lint, load_config
from xllib.lint.registry import registered_rules
from xllib.recalc import RecalcFailed, recalculate


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="xllib")
    commands = parser.add_subparsers(dest="command", required=True)
    lint_parser = commands.add_parser("lint", help="inspect and lint an .xlsx file")
    lint_parser.add_argument("path", type=Path)
    lint_parser.add_argument("--json", action="store_true", dest="as_json")
    lint_parser.add_argument("--measure", action="store_true")
    lint_parser.add_argument("--config", type=Path)
    return parser


def _measure(workbook: Any) -> dict[str, Any]:
    shapes: set[str] = set()
    formula_cells = 0
    max_depth = max_calls = 0
    defined_names = len(workbook.defined_names)
    sections = {}
    for sheet in workbook.sheets:
        sections[sheet.name] = estimate_section_count(sheet)
        for cell in sheet.cells:
            if cell.formula is None:
                continue
            formula_cells += 1
            # Shapes used to be collected from inferred runs, which need two
            # adjacent formula cells in one row. Every isolated formula was
            # therefore invisible: a workbook holding seven distinct formulas
            # measured 2, and one holding a single formula measured 0. The
            # budget of 40 was being calibrated against a proxy that missed
            # most of the file, so the calibration would have set the wrong
            # number and then failed builds against it.
            shapes.add(normalize_formula_shape(cell.formula, row=cell.row, column=cell.column))
            walk = walk_formula(cell.formula, workbook=workbook, sheet_index=sheet.index)
            max_depth = max(max_depth, walk.max_function_depth)
            max_calls = max(max_calls, walk.function_count)
    return {
        "sheets": len(workbook.sheets),
        "sheet_states": {sheet.name: sheet.state for sheet in workbook.sheets},
        "estimated_sections": sections,
        # Reported so the shape count has a visible denominator. A distinct
        # count means nothing without knowing how many formulas produced it.
        "formula_cells": formula_cells,
        "distinct_formula_shapes": len(shapes),
        "max_formula_depth": max_depth,
        "max_function_calls_per_cell": max_calls,
        "defined_names": defined_names,
    }


def _run_lint(args: argparse.Namespace) -> int:
    # Before anything else copies or opens the file. Recalculation hands the
    # target to backends ending in Excel COM, so a legacy `.xls` or a corrupt
    # archive would be launched in Excel despite openpyxl already knowing it
    # cannot be read — which ended the process with an access violation rather
    # than exit 3. Confirmed by running it; see `tests/test_cli.py`.
    assert_readable(args.path)
    if args.measure:
        workbook = load_workbook(args.path)
        print(json.dumps(_measure(workbook), indent=2))
        return 0
    config_path = args.config if args.config is not None else discover_config()
    config = load_config(config_path)
    try:
        with recalculate(args.path) as result:
            workbook = load_workbook(result.path)
            report = lint(workbook, config, registered_rules())
    except RecalcFailed:
        workbook = load_workbook(args.path)
        report = lint(workbook, config, registered_rules())
    print(report.to_json() if args.as_json else report.to_text())
    return report.exit_code


#: Exit 3 means the tool could not read the file, not that the workbook failed
#: lint. `BadZipFile` and `InvalidFileException` both inherit from `Exception`
#: and not from `OSError`, so the original `(OSError, ValueError)` pair let a
#: truncated `.xlsx` and a legacy `.xls` reach the user as a traceback — both
#: confirmed by running them, not assumed. Deliberately an explicit list and
#: not a bare `except Exception`: a genuine bug inside a rule must keep failing
#: loudly rather than be reported as an unreadable file.
UNREADABLE = (OSError, ValueError, BadZipFile, InvalidFileException)


def main(argv: Sequence[str] | None = None) -> int:
    try:
        args = _parser().parse_args(argv)
        if args.command == "lint":
            return _run_lint(args)
        return 3
    except UNREADABLE as error:
        print(f"xllib: {error}", file=sys.stderr)
        return 3


if __name__ == "__main__":
    raise SystemExit(main())
