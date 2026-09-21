"""The explicit, reproducible Phase 0 rule registry."""

from __future__ import annotations

from collections.abc import Iterable
from typing import Any

from openpyxl.utils import get_column_letter

from xllib.inspect import (
    Capability,
    Workbook,
    build_reference_graph,
    detect_period_axis,
    estimate_section_count,
    infer_formula_runs,
    walk_formula,
)

from ..rule import (
    Confidence,
    Finding,
    Locus,
    Rule,
    RuleContext,
    Severity,
    Status,
)


def _locus(workbook: Workbook, sheet_name: str, ref: str | None = None) -> Locus:
    sheet = workbook.sheet(sheet_name)
    row = column = 0
    if ref:
        cell = sheet.cell(1, 1) if sheet else None
        if sheet:
            cell = next((item for item in sheet.cells if item.coordinate == ref), cell)
        if cell and cell.coordinate == ref:
            row, column = cell.row, cell.column
    return Locus(
        sheet_name,
        ref,
        sheet.index if sheet else -1,
        row,
        column,
        sheet.state if sheet else None,
    )


def _finding(
    rule_id: str,
    status: Status,
    severity: Severity,
    confidence: Confidence,
    message: str,
    locus: Locus,
    remediation: str,
    evidence: dict[str, str] | None = None,
) -> Finding:
    return Finding(
        rule_id,
        status,
        severity,
        confidence,
        message,
        locus,
        evidence or {},
        remediation,
    )


def _violation(
    rule_id: str,
    confidence: Confidence,
    message: str,
    locus: Locus,
    remediation: str,
    evidence: dict[str, str] | None = None,
) -> Finding:
    return _finding(
        rule_id,
        Status.VIOLATION,
        Severity.ERROR,
        confidence,
        message,
        locus,
        remediation,
        evidence,
    )


def _xl001(workbook: Workbook, _: RuleContext) -> Iterable[Finding]:
    graph = build_reference_graph(workbook)
    for name in workbook.defined_names:
        if len(name.references) != 1 or not name.references[0].is_single_cell:
            continue
        target = name.references[0]
        sheet = workbook.sheet(target.sheet or "")
        if sheet is None or target.min_row is None or target.min_column is None:
            continue
        cell = sheet.cell(target.min_row, target.min_column)
        if cell is None or cell.formula is not None or not isinstance(cell.value, (int, float)):
            continue
        used = [
            edge
            for edge in graph.edges
            if edge.via_defined_name is not None
            and edge.via_defined_name.casefold() == name.name.casefold()
        ]
        left = sheet.cell(cell.row, cell.column - 1) if cell.column > 1 else None
        above = sheet.cell(cell.row - 1, cell.column) if cell.row > 1 else None
        labelled = any(
            isinstance(candidate.effective_value, str) and candidate.effective_value.strip()
            for candidate in (left, above)
            if candidate is not None
        )
        if used and not labelled:
            users = ", ".join(f"{edge.source_sheet}!{edge.source_coordinate}" for edge in used)
            yield _violation(
                "XL001",
                Confidence.CERTAIN,
                f"named constant {name.name} has no adjacent label; referenced by {users}",
                _locus(workbook, sheet.name, cell.coordinate),
                "add a text label immediately left of or above the named constant",
                {"defined_name": name.name, "referenced_by": users},
            )


def _xl002(workbook: Workbook, context: RuleContext) -> Iterable[Finding]:
    allowed = {
        float(value) for value in context.options.get("allowed_literals", [0, 1, -1, 12, 100])
    }
    exemptions: dict[str, list[int]] = context.options.get("positional_exemptions", {})
    for sheet in workbook.sheets:
        for cell in sheet.cells:
            if cell.formula is None:
                continue
            for literal in walk_formula(
                cell.formula, workbook=workbook, sheet_index=sheet.index
            ).numeric_literals:
                position = None if literal.argument_index is None else literal.argument_index + 1
                exempt_positions = exemptions.get(literal.function or "", [])
                if literal.value in allowed or position in exempt_positions:
                    continue
                yield _violation(
                    "XL002",
                    Confidence.CERTAIN,
                    f"numeric literal {literal.text} in formula",
                    _locus(workbook, sheet.name, cell.coordinate),
                    "declare the value as a labelled input and reference it",
                    {"formula": cell.formula, "literal": literal.text},
                )


def _xl003(workbook: Workbook, _: RuleContext) -> Iterable[Finding]:
    for sheet in workbook.sheets:
        for run in infer_formula_runs(sheet):
            shapes = tuple(dict.fromkeys(run.shapes))
            if len(shapes) <= 1:
                continue
            start = get_column_letter(run.start_column)
            end = get_column_letter(run.end_column)
            ref = f"{start}{run.row}:{end}{run.row}"
            yield _violation(
                "XL003",
                Confidence.HEURISTIC,
                f"{len(shapes)} formula shapes in inferred run {ref}",
                _locus(workbook, sheet.name, ref),
                "use one copied formula shape or separate the exceptional column",
                {"inferred_run": ref, "shapes": " | ".join(shapes)},
            )


def _numeric_format(number_format: str) -> bool:
    cleaned = number_format.lower().replace('"', "")
    return cleaned not in ("", "general", "@") and any(
        token in cleaned for token in ("0", "#", "$", "£", "€", "%")
    )


def _xl004(workbook: Workbook, _: RuleContext) -> Iterable[Finding]:
    for sheet in workbook.sheets:
        for cell in sheet.cells:
            if not isinstance(cell.effective_value, str):
                continue
            if not _numeric_format(cell.number_format):
                continue
            error = cell.is_error
            yield _violation(
                "XL004",
                Confidence.CERTAIN,
                f"cached error {cell.effective_value} in a numeric-formatted cell"
                if error
                else "text value in a numeric-formatted cell",
                _locus(workbook, sheet.name, cell.coordinate),
                "repair the formula, or handle the error case so the cell holds a number"
                if error
                else "move status text to a status column and keep the numeric cell numeric",
                {
                    "value": cell.effective_value,
                    "number_format": cell.number_format,
                    "kind": "error" if error else "text",
                },
            )


def _xl005(workbook: Workbook, _: RuleContext) -> Iterable[Finding]:
    graph = build_reference_graph(workbook)
    for aggregate in graph.aggregate_ranges:
        target = aggregate.target
        sheet = workbook.sheet(target.sheet or "")
        bounds = (target.min_row, target.min_column, target.max_row, target.max_column)
        if sheet is None or any(value is None for value in bounds):
            continue
        min_row, min_col, max_row, max_col = bounds
        assert min_row is not None and min_col is not None
        assert max_row is not None and max_col is not None
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                cell = sheet.cell(row, column)
                if cell is None or not isinstance(cell.effective_value, str):
                    continue
                source = f"{aggregate.source_sheet}!{aggregate.source_coordinate}"
                error = cell.is_error
                yield _violation(
                    "XL005",
                    Confidence.CERTAIN,
                    f"cached error {cell.effective_value} inside range aggregated by {source}"
                    if error
                    else f"text value inside range aggregated by {source}",
                    _locus(workbook, sheet.name, cell.coordinate),
                    "repair the formula that errors; the aggregate cannot evaluate around it"
                    if error
                    else "remove text from the aggregate range or use a separate status column",
                    {
                        "aggregate": source,
                        "function": aggregate.function,
                        "kind": "error" if error else "text",
                    },
                )


def _xl006(workbook: Workbook, _: RuleContext) -> Iterable[Finding]:
    axes = [axis for sheet in workbook.sheets if (axis := detect_period_axis(sheet))]
    if not axes:
        yield _finding(
            "XL006",
            Status.INFO,
            Severity.ERROR,
            Confidence.HEURISTIC,
            "no period axis detected; inconsistent-period-axis was not evaluated",
            Locus(),
            "add a detectable period header or ignore this informational finding",
        )
        return
    for index, left in enumerate(axes):
        for right in axes[index + 1 :]:
            if left.labels != right.labels or left.start_column == right.start_column:
                continue
            yield _violation(
                "XL006",
                Confidence.HEURISTIC,
                f"same period axis starts in column {left.start_column} on {left.sheet} "
                f"and {right.start_column} on {right.sheet}",
                Locus(),
                "align matching period sequences to the same column index",
                {"left": left.sheet, "right": right.sheet},
            )


def _xl101(workbook: Workbook, context: RuleContext) -> Iterable[Finding]:
    limit = context.thresholds.get("sections_per_sheet", 5)
    for sheet in workbook.sheets:
        count = estimate_section_count(sheet)
        if count > limit:
            yield _violation(
                "XL101",
                Confidence.HEURISTIC,
                f"estimated {count} sections exceeds budget {limit}",
                _locus(workbook, sheet.name),
                "consolidate sections or document why the estimate is misleading",
                {"estimated_sections": str(count), "budget": str(limit)},
            )


def _check_names(workbook: Workbook) -> tuple[Any, ...]:
    return tuple(name for name in workbook.defined_names if name.name.casefold().startswith("chk_"))


def _xl102(workbook: Workbook, _: RuleContext) -> Iterable[Finding]:
    if not _check_names(workbook):
        yield _violation(
            "XL102",
            Confidence.CERTAIN,
            "no chk_* defined name found",
            Locus(),
            "define at least one check cell using the chk_* naming convention",
        )


def _xl103(workbook: Workbook, _: RuleContext) -> Iterable[Finding]:
    for name in _check_names(workbook):
        for target in name.references:
            sheet = workbook.sheet(target.sheet or "")
            if sheet is None or target.min_row is None or target.min_column is None:
                continue
            cell = sheet.cell(target.min_row, target.min_column)
            value = cell.effective_value if cell else None
            if value is True or value == 0:
                continue
            ref = cell.coordinate if cell else target.text
            yield _violation(
                "XL103",
                Confidence.CERTAIN,
                f"check {name.name} evaluates to {value!r}",
                _locus(workbook, sheet.name, ref),
                "correct the model until the check evaluates to 0 or TRUE",
                {"defined_name": name.name, "value": repr(value)},
            )


def _xl104(workbook: Workbook, _: RuleContext) -> Iterable[Finding]:
    graph = build_reference_graph(workbook)
    seen: set[tuple[str, int, int]] = set()
    for edge in graph.edges:
        target = edge.target
        sheet = workbook.sheet(target.sheet or "")
        if sheet is None or not target.is_single_cell:
            continue
        assert target.min_row is not None and target.min_column is not None
        cell = sheet.cell(target.min_row, target.min_column)
        if cell is None or cell.formula is not None or cell.value is None:
            continue
        if cell.has_font and cell.has_fill:
            continue
        key = (sheet.name, cell.row, cell.column)
        if key in seen:
            continue
        seen.add(key)
        missing = []
        if not cell.has_font:
            missing.append("font")
        if not cell.has_fill:
            missing.append("fill")
        yield _violation(
            "XL104",
            Confidence.HEURISTIC,
            "formula-dependent constant has no distinct style",
            _locus(workbook, sheet.name, cell.coordinate),
            "apply redundant input styling (font and fill) and provide an on-sheet legend",
            {"missing": ", ".join(missing)},
        )


def _rule(
    rule_id: str,
    slug: str,
    severity: Severity,
    confidence: Confidence,
    requires: frozenset[Capability],
    check: Any,
) -> Rule:
    return Rule(
        rule_id,
        slug,
        slug.replace("-", " ").title(),
        severity,
        confidence,
        requires,
        {},
        check,
        f"docs/rules/{rule_id}.md",
    )


RULES = (
    _rule(
        "XL001",
        "unlabelled-named-constant",
        Severity.ERROR,
        Confidence.CERTAIN,
        frozenset({Capability.REF_GRAPH}),
        _xl001,
    ),
    _rule(
        "XL002",
        "numeric-literal-in-formula",
        Severity.ERROR,
        Confidence.CERTAIN,
        frozenset({Capability.FORMULA_TOKENS}),
        _xl002,
    ),
    _rule(
        "XL003",
        "mixed-formula-shape-in-run",
        Severity.ERROR,
        Confidence.HEURISTIC,
        frozenset({Capability.FORMULA_TOKENS}),
        _xl003,
    ),
    _rule(
        "XL004",
        "text-in-numeric-cell",
        Severity.ERROR,
        Confidence.CERTAIN,
        frozenset({Capability.CACHED_VALUES, Capability.NUMBER_FORMATS}),
        _xl004,
    ),
    _rule(
        "XL005",
        "text-in-aggregated-range",
        Severity.ERROR,
        Confidence.CERTAIN,
        frozenset({Capability.CACHED_VALUES, Capability.REF_GRAPH}),
        _xl005,
    ),
    _rule(
        "XL006",
        "inconsistent-period-axis",
        Severity.ERROR,
        Confidence.HEURISTIC,
        frozenset(),
        _xl006,
    ),
    _rule(
        "XL101", "estimated-section-count", Severity.WARN, Confidence.HEURISTIC, frozenset(), _xl101
    ),
    _rule(
        "XL102", "no-check-convention-found", Severity.WARN, Confidence.CERTAIN, frozenset(), _xl102
    ),
    _rule(
        "XL103",
        "check-not-passing",
        Severity.WARN,
        Confidence.CERTAIN,
        frozenset({Capability.CACHED_VALUES}),
        _xl103,
    ),
    _rule(
        "XL104",
        "input-styling-not-redundant",
        Severity.WARN,
        Confidence.HEURISTIC,
        frozenset({Capability.REF_GRAPH}),
        _xl104,
    ),
)

__all__ = ["RULES"]
