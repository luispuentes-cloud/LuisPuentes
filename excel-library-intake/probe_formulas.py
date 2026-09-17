"""Does `formulas` actually recalculate our fixture shapes on this Python?

The CI path depends on the answer: Linux CI has no Excel, so this is the only
backend that can evaluate value-dependent lint rules there.
"""
import sys
import tempfile
from pathlib import Path

print("python", sys.version.split()[0])

try:
    import formulas
    print("formulas", getattr(formulas, "__version__", "?"))
except Exception as exc:
    print("IMPORT FAILED:", type(exc).__name__, exc)
    raise SystemExit(1)

import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Calc"
# Shapes the fixtures rely on: aggregate, guard, named-range driver, check cell.
for i, v in enumerate([10, 20, 30, 40], start=4):
    ws.cell(row=i, column=9, value=v)      # I4:I7 load
    ws.cell(row=i, column=11, value=v * 2)  # K4:K7 base
ws["C1"] = 0.85
wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName("FTE_TARGET", attr_text="Calc!$C$1"))
ws["D4"] = "=SUM($I$4:$I$7)"
ws["D5"] = "=K4*FTE_TARGET/SUM($I$4:$I$7)"
ws["D6"] = "=IF(SUM($I$4:$I$7)=0,\"N/A\",K4/SUM($I$4:$I$7))"
ws["D7"] = "=ROUND(D5,2)"
ws["D8"] = "=SUM(D4:D7)-SUM(D4:D7)"  # a chk_ style control, must be 0

path = Path(tempfile.gettempdir()) / "xllib_spike.xlsx"
wb.save(path)
print("wrote", path)

pre = openpyxl.load_workbook(path, data_only=True)
print("cached values before recalc:", [pre["Calc"][c].value for c in ("D4", "D5", "D8")])

try:
    xl = formulas.ExcelModel().loads(str(path)).finish()
    sol = xl.calculate()
    hits = {k: v for k, v in sol.items() if "D4" in k or "D5" in k or "D8" in k}
    print("solution entries returned:", len(sol))
    for k, v in list(hits.items())[:8]:
        print("  ", k, "=", v)
except Exception as exc:
    print("CALCULATE FAILED:", type(exc).__name__, exc)
    raise SystemExit(2)

out = Path(tempfile.gettempdir()) / "xllib_spike_out"
try:
    xl.write(dirpath=str(out))
    print("write-back OK ->", out)
except Exception as exc:
    print("WRITE-BACK FAILED:", type(exc).__name__, exc)
